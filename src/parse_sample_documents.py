from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def safe_name(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", value)


def extract_document(path: Path) -> tuple[dict, str]:
    document = Document(path)
    paragraphs: list[dict] = []
    heading_count = 0
    replacement_count = 0

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        is_heading = style.lower().startswith("heading") or style.startswith("标题")
        heading_count += int(is_heading)
        replacement_count += text.count("\ufffd")
        paragraphs.append({"index": index, "style": style, "text": text, "is_heading": is_heading})

    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            replacement_count += sum(cell.count("\ufffd") for cell in cells)
            rows.append(cells)
        tables.append(rows)

    image_count = sum(1 for relation in document.part.rels.values() if "image" in relation.reltype)
    total_chars = sum(len(item["text"]) for item in paragraphs) + sum(
        len(cell) for table in tables for row in table for cell in row
    )
    text_ok = total_chars >= 500 and len(paragraphs) > 0
    encoding_ok = replacement_count == 0
    heading_status = "章节标题可追溯" if heading_count > 0 else "未识别标题样式，仅可用段落序号"
    table_status = "已提取" if tables else "未检测到表格，需人工核对原文"
    overall = "通过" if text_ok and encoding_ok and heading_count > 0 else "部分通过"
    if not text_ok or not encoding_ok:
        overall = "失败"

    stats = {
        "paragraph_count": len(paragraphs),
        "heading_count": heading_count,
        "table_count": len(tables),
        "image_count": image_count,
        "total_chars": total_chars,
        "replacement_char_count": replacement_count,
        "text_status": "通过" if text_ok else "失败",
        "encoding_status": "通过" if encoding_ok else "失败",
        "heading_status": heading_status,
        "table_status": table_status,
        "overall_status": overall,
    }

    output: list[str] = [
        f"# {path.stem}",
        "",
        f"来源文件：`{path.name}`",
        "",
        "## 正文",
        "",
    ]
    for item in paragraphs:
        if item["is_heading"]:
            output.append(f"### {item['text']}")
        else:
            output.append(f"[{item['index']}] {item['text']}")
        output.append("")
    if tables:
        output.extend(["## 表格", ""])
        for table_index, table in enumerate(tables, start=1):
            output.append(f"### 表格 {table_index}")
            output.append("")
            for row_index, row in enumerate(table, start=1):
                output.append(f"- 行 {row_index}: " + " | ".join(cell.replace("|", "\\|") for cell in row))
            output.append("")
    return stats, "\n".join(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("project_root", type=Path)
    args = parser.parse_args()

    root = args.project_root.resolve()
    manifest = json.loads((root / "data" / "corpus_manifest.json").read_text(encoding="utf-8"))
    preparation = json.loads((root / "_staging" / "20260907_语料基线" / "word_preparation.json").read_text(encoding="utf-8-sig"))
    prepared_by_id = {item["source_id"]: item for item in preparation}
    output_root = root / "data" / "processed" / "samples"
    output_root.mkdir(parents=True, exist_ok=True)
    results: list[dict] = []

    samples = [item for item in manifest["documents"] if item["sample_selected"]]
    if len(samples) != 6:
        raise RuntimeError(f"抽检样本数量应为 6，实际为 {len(samples)}")

    for item in samples:
        prepared = prepared_by_id[item["source_id"]]
        normalized = root / Path(prepared["normalized_docx"])
        stats, markdown = extract_document(normalized)
        output_path = output_root / f"{safe_name(Path(item['file_name']).stem)}.md"
        output_path.write_text(markdown, encoding="utf-8")
        results.append(
            {
                "source_id": item["source_id"],
                "product_group": item["product_group"],
                "document_type": item["document_type"],
                "file_name": item["file_name"],
                "original_extension": item["extension"],
                "conversion_status": prepared["conversion_status"],
                "pdf_status": prepared["pdf_status"],
                "processed_relative_path": output_path.relative_to(root).as_posix(),
                **stats,
            }
        )

    result_path = root / "reports" / "parse_check_data.json"
    result_path.parent.mkdir(parents=True, exist_ok=True)
    result_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "result": str(result_path),
        "sample_count": len(results),
        "passed": sum(1 for item in results if item["overall_status"] == "通过"),
        "partial": sum(1 for item in results if item["overall_status"] == "部分通过"),
        "failed": sum(1 for item in results if item["overall_status"] == "失败"),
    }, ensure_ascii=True))


if __name__ == "__main__":
    main()
